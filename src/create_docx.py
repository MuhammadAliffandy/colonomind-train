import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_newline_to_cell(cell, text):
    """
    Format string like "70.35% (63.67%–76.26%)" 
    into two lines in a docx cell.
    """
    if " (" in text:
        parts = text.split(" (")
        cell.text = parts[0]
        p = cell.paragraphs[0]
        run = p.add_run(f"\n({parts[1]}")
    else:
        cell.text = text

def main():
    results_dir = "Manuscript_Consistent_Results"
    output_file = os.path.join(results_dir, "Final_Results_Manuscript.docx")
    
    doc = Document()
    
    # Title
    title = doc.add_heading('ColonoMind Final Consistent Results', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    datasets = ['NTUH', 'TMC-UCM', 'LIMUC', 'Unified']
    
    doc.add_heading('1. Performance Tables', level=1)
    
    for i, dataset in enumerate(datasets, 1):
        doc.add_heading(f'Table {i}: Performance on {dataset} Dataset', level=2)
        
        csv_path = os.path.join(results_dir, f"Table_{i}_{dataset}.csv")
        if not os.path.exists(csv_path):
            continue
            
        with open(csv_path, 'r', encoding='utf-8') as f:
            lines = f.readlines()
            
        # Create table
        table = doc.add_table(rows=1, cols=6)
        table.style = 'Table Grid'
        
        # Header
        hdr_cells = table.rows[0].cells
        headers = ['Model', 'Accuracy (%)', 'Precision (%)', 'Recall (%)', 'F1 (%)', 'QWK']
        for j, text in enumerate(headers):
            hdr_cells[j].text = text
            # Make header bold
            for run in hdr_cells[j].paragraphs[0].runs:
                run.font.bold = True
                
        # Data rows
        for line in lines[1:]:
            cols = line.strip().split(',')
            if len(cols) < 6:
                continue
                
            row_cells = table.add_row().cells
            # Order in CSV: Model, Accuracy, F1, Precision, Recall, QWK
            # We want: Model, Accuracy, Precision, Recall, F1, QWK
            model = cols[0]
            acc = cols[1]
            f1 = cols[2]
            prec = cols[3]
            rec = cols[4]
            qwk = cols[5]
            
            ordered_cols = [model, acc, prec, rec, f1, qwk]
            
            for j, text in enumerate(ordered_cols):
                add_newline_to_cell(row_cells[j], text)
                
        doc.add_paragraph('\n')
        
    # Table 5: Ensemble Voting
    doc.add_heading('Table 5: Ensemble Voting', level=2)
    csv_path = os.path.join(results_dir, "Table_5_Ensemble_Voting.csv")
    if os.path.exists(csv_path):
        with open(csv_path, 'r', encoding='utf-8') as f:
            lines = f.readlines()
        table = doc.add_table(rows=1, cols=7)
        table.style = 'Table Grid'
        hdr_cells = table.rows[0].cells
        headers = ['Dataset', 'Model', 'Accuracy (%)', 'Precision (%)', 'Recall (%)', 'F1 (%)', 'QWK']
        for j, text in enumerate(headers):
            hdr_cells[j].text = text
            for run in hdr_cells[j].paragraphs[0].runs: run.font.bold = True
        for line in lines[1:]:
            cols = line.strip().split(',')
            if len(cols) < 7: continue
            row_cells = table.add_row().cells
            ds, model, acc, prec, rec, f1, qwk = cols[0], cols[1], cols[2], cols[3], cols[4], cols[5], cols[6]
            ordered_cols = [ds, model, acc, prec, rec, f1, qwk]
            for j, text in enumerate(ordered_cols):
                add_newline_to_cell(row_cells[j], text)
    doc.add_paragraph('\n')
    
    # Table 6: Weighted Voting
    doc.add_heading('Table 6: Weighted Voting', level=2)
    csv_path = os.path.join(results_dir, "Table_6_Weighted_Voting.csv")
    if os.path.exists(csv_path):
        with open(csv_path, 'r', encoding='utf-8') as f:
            lines = f.readlines()
        table = doc.add_table(rows=1, cols=7)
        table.style = 'Table Grid'
        hdr_cells = table.rows[0].cells
        headers = ['Dataset', 'Model', 'Accuracy (%)', 'Precision (%)', 'Recall (%)', 'F1 (%)', 'QWK']
        for j, text in enumerate(headers):
            hdr_cells[j].text = text
            for run in hdr_cells[j].paragraphs[0].runs: run.font.bold = True
        for line in lines[1:]:
            cols = line.strip().split(',')
            if len(cols) < 7: continue
            row_cells = table.add_row().cells
            ds, model, acc, prec, rec, f1, qwk = cols[0], cols[1], cols[2], cols[3], cols[4], cols[5], cols[6]
            ordered_cols = [ds, model, acc, prec, rec, f1, qwk]
            for j, text in enumerate(ordered_cols):
                add_newline_to_cell(row_cells[j], text)
    doc.add_paragraph('\n')
        
    doc.add_page_break()
    
    # Confusion Matrices
    doc.add_heading('2. Confusion Matrices', level=1)
    for dataset in datasets:
        doc.add_heading(f'Confusion Matrices - {dataset}', level=2)
        img_path = os.path.join(results_dir, f"Fig_CM_{dataset}.png")
        if os.path.exists(img_path):
            doc.add_picture(img_path, width=Inches(6.0))
        doc.add_paragraph('\n')
        
    doc.add_page_break()
    
    # ROC Curves
    doc.add_heading('3. ROC Curves', level=1)
    img_path = os.path.join(results_dir, "Fig_ROC_Combined.png")
    if os.path.exists(img_path):
        doc.add_picture(img_path, width=Inches(6.5))
        
    doc.add_page_break()
    
    # Narration Snippet
    doc.add_heading('4. AUC Narration Snippet', level=1)
    txt_path = os.path.join(results_dir, "narration_snippet.txt")
    if os.path.exists(txt_path):
        with open(txt_path, 'r', encoding='utf-8') as f:
            text = f.read()
        p = doc.add_paragraph(text)
        
    doc.save(output_file)
    print(f"✅ Generated {output_file}")

if __name__ == "__main__":
    main()
